"""Fase B - Excel para SQLite.

openpyxl em read_only, sqlite3 da biblioteca padrão, esquema explícito para as
nove tabelas, seis chaves estrangeiras declaradas, PRAGMA foreign_keys=ON, pais
antes de filhos, uma transação por aba, lotes de 5.000 linhas, datas em ISO 8601,
nulos preservados e índices criados depois da carga. Sem pandas e sem números de
resultado fixos: o esperado vem de contrato/manifesto_olist.json.

Grava saida/olist.sqlite e saida/carga_sqlite.json. Sai com código 1 em qualquer
divergência de contagem, unicidade, integridade, chave estrangeira ou nulos.
"""
import hashlib
import json
import sqlite3
import sys
import time
from datetime import date, datetime
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
ENTRADA = BASE / "entrada"
SAIDA = BASE / "saida"
MANIFESTO = BASE / "contrato" / "manifesto_olist.json"
PERFIL = SAIDA / "perfil_excel.json"
DESTINO = SAIDA / "olist.sqlite"
LOTE = 5000
LIMITE_NOME_ABA = 31

# Tipos lógicos: T = TEXT, I = INTEGER, R = REAL, D = data/hora normalizada como TEXT ISO 8601.
ESQUEMA = {
    "olist_customers": {
        "colunas": [("customer_id", "T"), ("customer_unique_id", "T"), ("customer_zip_code_prefix", "I"),
                    ("customer_city", "T"), ("customer_state", "T")],
        "pk": ["customer_id"], "fks": [], "not_null": ["customer_unique_id"],
    },
    "olist_sellers": {
        "colunas": [("seller_id", "T"), ("seller_zip_code_prefix", "I"), ("seller_city", "T"),
                    ("seller_state", "T")],
        "pk": ["seller_id"], "fks": [], "not_null": [],
    },
    "olist_products": {
        "colunas": [("product_id", "T"), ("product_category_name", "T"), ("product_name_lenght", "I"),
                    ("product_description_lenght", "I"), ("product_photos_qty", "I"),
                    ("product_weight_g", "I"), ("product_length_cm", "I"), ("product_height_cm", "I"),
                    ("product_width_cm", "I")],
        "pk": ["product_id"], "fks": [], "not_null": [],
    },
    "product_category_name_translation": {
        "colunas": [("product_category_name", "T"), ("product_category_name_english", "T")],
        "pk": ["product_category_name"], "fks": [], "not_null": [],
    },
    "olist_geolocation": {
        "colunas": [("geolocation_zip_code_prefix", "I"), ("geolocation_lat", "R"), ("geolocation_lng", "R"),
                    ("geolocation_city", "T"), ("geolocation_state", "T")],
        "pk": [], "fks": [], "not_null": [],
    },
    "olist_orders": {
        "colunas": [("order_id", "T"), ("customer_id", "T"), ("order_status", "T"),
                    ("order_purchase_timestamp", "D"), ("order_approved_at", "D"),
                    ("order_delivered_carrier_date", "D"), ("order_delivered_customer_date", "D"),
                    ("order_estimated_delivery_date", "D")],
        "pk": ["order_id"], "fks": [("customer_id", "olist_customers", "customer_id")],
        "not_null": ["customer_id"],
    },
    "olist_order_items": {
        "colunas": [("order_id", "T"), ("order_item_id", "I"), ("product_id", "T"), ("seller_id", "T"),
                    ("shipping_limit_date", "D"), ("price", "R"), ("freight_value", "R")],
        "pk": ["order_id", "order_item_id"],
        "fks": [("order_id", "olist_orders", "order_id"), ("product_id", "olist_products", "product_id"),
                ("seller_id", "olist_sellers", "seller_id")],
        "not_null": ["product_id", "seller_id"],
    },
    "olist_order_payments": {
        "colunas": [("order_id", "T"), ("payment_sequential", "I"), ("payment_type", "T"),
                    ("payment_installments", "I"), ("payment_value", "R")],
        "pk": ["order_id", "payment_sequential"], "fks": [("order_id", "olist_orders", "order_id")],
        "not_null": [],
    },
    "olist_order_reviews": {
        "colunas": [("review_id", "T"), ("order_id", "T"), ("review_score", "I"), ("review_comment_title", "T"),
                    ("review_comment_message", "T"), ("review_creation_date", "D"),
                    ("review_answer_timestamp", "D")],
        "pk": ["review_id", "order_id"], "fks": [("order_id", "olist_orders", "order_id")],
        "not_null": [],
    },
}
# Pais antes de filhos.
ORDEM_CARGA = ["olist_customers", "olist_sellers", "olist_products", "product_category_name_translation",
               "olist_geolocation", "olist_orders", "olist_order_items", "olist_order_payments",
               "olist_order_reviews"]
INDICES = [
    ("idx_orders_customer", "olist_orders", ["customer_id"]),
    ("idx_orders_purchase", "olist_orders", ["order_purchase_timestamp"]),
    ("idx_orders_status", "olist_orders", ["order_status"]),
    ("idx_items_product", "olist_order_items", ["product_id"]),
    ("idx_items_seller", "olist_order_items", ["seller_id"]),
    ("idx_reviews_order", "olist_order_reviews", ["order_id"]),
    ("idx_customers_state", "olist_customers", ["customer_state"]),
    ("idx_products_category", "olist_products", ["product_category_name"]),
    ("idx_geolocation_prefix", "olist_geolocation", ["geolocation_zip_code_prefix"]),
]
TIPO_SQL = {"T": "TEXT", "I": "INTEGER", "R": "REAL", "D": "TEXT"}


class ErroConversao(Exception):
    pass


def converter(valor, tipo, conversoes):
    if valor is None:
        return None
    if tipo == "T":
        if isinstance(valor, str):
            return valor
        conversoes["texto_de_outro_tipo"] += 1
        return str(valor)
    if tipo == "I":
        if isinstance(valor, bool):
            raise ErroConversao(f"booleano {valor!r} em coluna inteira")
        if isinstance(valor, int):
            return valor
        if isinstance(valor, float) and valor.is_integer():
            conversoes["inteiro_de_float"] += 1
            return int(valor)
        if isinstance(valor, str) and valor.strip().lstrip("-").isdigit():
            conversoes["inteiro_de_texto"] += 1
            return int(valor.strip())
        raise ErroConversao(f"valor {valor!r} não é inteiro")
    if tipo == "R":
        if isinstance(valor, (int, float)) and not isinstance(valor, bool):
            return float(valor)
        if isinstance(valor, str):
            conversoes["real_de_texto"] += 1
            return float(valor)
        raise ErroConversao(f"valor {valor!r} não é numérico")
    if tipo == "D":
        if isinstance(valor, datetime):
            dt = valor
        elif isinstance(valor, date):
            dt = datetime(valor.year, valor.month, valor.day)
        elif isinstance(valor, str):
            try:
                dt = datetime.fromisoformat(valor.strip())
            except ValueError as e:
                raise ErroConversao(f"data {valor!r} fora do padrão") from e
        else:
            raise ErroConversao(f"valor {valor!r} não é data")
        conversoes["data_normalizada_iso8601"] += 1
        return dt.isoformat(timespec="seconds")
    raise ErroConversao(f"tipo lógico desconhecido {tipo}")


def sha256(caminho):
    h = hashlib.sha256()
    with open(caminho, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()


def ddl(tabela):
    e = ESQUEMA[tabela]
    partes = []
    for nome, tipo in e["colunas"]:
        restricao = " NOT NULL" if nome in e["not_null"] or (nome in e["pk"]) else ""
        partes.append(f"  {nome} {TIPO_SQL[tipo]}{restricao}")
    if e["pk"]:
        partes.append(f"  PRIMARY KEY ({', '.join(e['pk'])})")
    for col, pai, col_pai in e["fks"]:
        partes.append(f"  FOREIGN KEY ({col}) REFERENCES {pai}({col_pai})")
    return f"CREATE TABLE {tabela} (\n" + ",\n".join(partes) + "\n)"


def localizar_aba(wb, tabela):
    if tabela in wb.sheetnames:
        return tabela
    if tabela[:LIMITE_NOME_ABA] in wb.sheetnames:
        return tabela[:LIMITE_NOME_ABA]
    raise SystemExit(f"FALHA: aba da tabela {tabela} não encontrada")


def carregar_tabela(con, tabela, arquivo):
    e = ESQUEMA[tabela]
    nomes = [c for c, _ in e["colunas"]]
    tipos = [t for _, t in e["colunas"]]
    conversoes = {"texto_de_outro_tipo": 0, "inteiro_de_float": 0, "inteiro_de_texto": 0,
                  "real_de_texto": 0, "data_normalizada_iso8601": 0}
    t0 = time.perf_counter()
    wb = openpyxl.load_workbook(ENTRADA / arquivo, read_only=True, data_only=True)
    aba = localizar_aba(wb, tabela)
    linhas = wb[aba].iter_rows(values_only=True)
    cabecalho = list(next(linhas))
    if cabecalho != nomes:
        wb.close()
        raise SystemExit(f"FALHA: cabeçalho de {tabela} difere do esquema explícito: {cabecalho}")
    sql = f"INSERT INTO {tabela} ({', '.join(nomes)}) VALUES ({', '.join('?' * len(nomes))})"
    lote, n, lotes = [], 0, 0
    con.execute("BEGIN")
    try:
        largura = len(nomes)
        for i, linha in enumerate(linhas, start=2):
            # read_only não completa células vazias no fim da linha: completa com nulo.
            if len(linha) < largura:
                linha = tuple(linha) + (None,) * (largura - len(linha))
            elif len(linha) > largura:
                if any(v is not None for v in linha[largura:]):
                    raise SystemExit(f"FALHA: {tabela}, linha {i}: célula preenchida além do cabeçalho")
                linha = linha[:largura]
            try:
                lote.append(tuple(converter(v, t, conversoes) for v, t in zip(linha, tipos)))
            except ErroConversao as err:
                raise SystemExit(f"FALHA de conversão em {tabela}, linha {i} da aba: {err}")
            if len(lote) == LOTE:
                con.executemany(sql, lote)
                n += len(lote)
                lotes += 1
                lote = []
        if lote:
            con.executemany(sql, lote)
            n += len(lote)
            lotes += 1
        con.execute("COMMIT")
    except sqlite3.IntegrityError as err:
        con.execute("ROLLBACK")
        raise SystemExit(f"FALHA de integridade ao carregar {tabela} (lote {lotes + 1}): {err}")
    finally:
        wb.close()
    return {"tabela": tabela, "arquivo": arquivo, "aba": aba, "linhas_inseridas": n, "lotes": lotes,
            "conversoes": conversoes, "segundos": round(time.perf_counter() - t0, 3)}


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    inicio = time.perf_counter()
    manifesto = json.loads(MANIFESTO.read_text(encoding="utf-8"))
    tabela_arquivo, esperado, chaves = {}, {}, {}
    for w in manifesto["workbooks"]:
        for s in w["sheets"]:
            tabela_arquivo[s["table"]] = w["file"]
            esperado[s["table"]] = s["rows"]
            chaves[s["table"]] = s["candidate_key"]
    if set(esperado) != set(ESQUEMA):
        raise SystemExit(f"FALHA: tabelas do manifesto {sorted(esperado)} diferem do esquema")

    hashes_antes = {w["file"]: sha256(ENTRADA / w["file"]) for w in manifesto["workbooks"]}
    for w in manifesto["workbooks"]:
        if hashes_antes[w["file"]] != w["sha256"]:
            raise SystemExit(f"FALHA: SHA256 de {w['file']} difere do manifesto")

    SAIDA.mkdir(exist_ok=True)
    if DESTINO.exists():
        DESTINO.unlink()
    con = sqlite3.connect(DESTINO, isolation_level=None)
    con.execute("PRAGMA foreign_keys = ON")
    fk_ativo = con.execute("PRAGMA foreign_keys").fetchone()[0]
    if fk_ativo != 1:
        raise SystemExit("FALHA: PRAGMA foreign_keys não ficou ativo")

    for tabela in ORDEM_CARGA:
        con.execute(ddl(tabela))

    cargas = []
    for tabela in ORDEM_CARGA:
        print(f"carregando {tabela} ...", flush=True)
        r = carregar_tabela(con, tabela, tabela_arquivo[tabela])
        print(f"  {r['linhas_inseridas']} linhas em {r['lotes']} lotes, {r['segundos']} s", flush=True)
        cargas.append(r)

    t_idx = time.perf_counter()
    for nome, tabela, cols in INDICES:
        con.execute(f"CREATE INDEX {nome} ON {tabela} ({', '.join(cols)})")
    con.execute("ANALYZE")
    seg_indices = round(time.perf_counter() - t_idx, 3)

    falhas = []
    # Gate 1: nove tabelas.
    tabelas_db = [r[0] for r in con.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]
    if sorted(tabelas_db) != sorted(ESQUEMA):
        falhas.append(f"tabelas no banco: {tabelas_db}")

    # Gate 2: contagens = manifesto; unicidade das chaves candidatas.
    contagens = []
    for tabela in ORDEM_CARGA:
        n = con.execute(f"SELECT count(*) FROM {tabela}").fetchone()[0]
        k = chaves[tabela]
        if k:
            distintos = con.execute(f"SELECT count(*) FROM (SELECT DISTINCT {', '.join(k)} FROM {tabela})").fetchone()[0]
        else:
            distintos = None
        ok = n == esperado[tabela] and (distintos is None or distintos == n)
        if not ok:
            falhas.append(f"{tabela}: {n} linhas (manifesto {esperado[tabela]}), chave distinta {distintos}")
        contagens.append({"tabela": tabela, "linhas_sqlite": n, "linhas_manifesto": esperado[tabela],
                          "chave_candidata": k, "chaves_distintas": distintos,
                          "chave_unica": None if distintos is None else distintos == n, "confere": ok})

    # Gate 3: integridade.
    integrity = [r[0] for r in con.execute("PRAGMA integrity_check")]
    if integrity != ["ok"]:
        falhas.append(f"integrity_check: {integrity[:5]}")

    # Gate 4: FKs declaradas e verificadas.
    fks_declaradas = []
    for tabela in ORDEM_CARGA:
        for r in con.execute(f"PRAGMA foreign_key_list('{tabela}')"):
            fks_declaradas.append({"tabela": tabela, "coluna": r[3], "referencia": f"{r[2]}({r[4]})"})
    fk_check = con.execute("PRAGMA foreign_key_check").fetchall()
    if fk_check:
        falhas.append(f"foreign_key_check retornou {len(fk_check)} linhas")
    orfaos = []
    for tabela in ORDEM_CARGA:
        for col, pai, col_pai in ESQUEMA[tabela]["fks"]:
            q = (f"SELECT count(*) FROM {tabela} f LEFT JOIN {pai} p ON p.{col_pai} = f.{col} "
                 f"WHERE f.{col} IS NOT NULL AND p.{col_pai} IS NULL")
            n_orf = con.execute(q).fetchone()[0]
            orfaos.append({"relacao": f"{tabela}.{col} -> {pai}.{col_pai}", "linhas_orfas": n_orf})
            if n_orf:
                falhas.append(f"{tabela}.{col}: {n_orf} linhas sem pai")
    n_fks_esperadas = sum(len(ESQUEMA[t]["fks"]) for t in ESQUEMA)
    if len(fks_declaradas) != n_fks_esperadas:
        falhas.append(f"{len(fks_declaradas)} FKs declaradas, esquema pede {n_fks_esperadas}")

    # Gate 5: nulos preservados (comparados ao perfil do Excel, quando existir).
    nulos = []
    perfil = json.loads(PERFIL.read_text(encoding="utf-8")) if PERFIL.exists() else None
    nulos_excel = {}
    if perfil:
        for a in perfil["arquivos"]:
            for s in a["abas"]:
                for c in s["colunas"]:
                    nulos_excel[(s["tabela"], c["nome"])] = c["tipos_observados"].get("null", 0)
    for tabela in ORDEM_CARGA:
        for col, _ in ESQUEMA[tabela]["colunas"]:
            n_nulo = con.execute(f"SELECT count(*) FROM {tabela} WHERE {col} IS NULL").fetchone()[0]
            n_ex = nulos_excel.get((tabela, col))
            if n_nulo or n_ex:
                nulos.append({"tabela": tabela, "coluna": col, "nulos_sqlite": n_nulo, "nulos_excel": n_ex})
            if n_ex is not None and n_ex != n_nulo:
                falhas.append(f"{tabela}.{col}: {n_nulo} nulos no SQLite, {n_ex} no Excel")

    con.close()
    hashes_depois = {w["file"]: sha256(ENTRADA / w["file"]) for w in manifesto["workbooks"]}
    entrada_intacta = hashes_antes == hashes_depois
    if not entrada_intacta:
        falhas.append("SHA256 da entrada mudou durante a carga")

    relatorio = {
        "gerado_em": datetime.now().astimezone().isoformat(timespec="seconds"),
        "script": "src/01_excel_para_sqlite.py",
        "banco": "saida/olist.sqlite",
        "bytes_banco": DESTINO.stat().st_size,
        "versoes": {"python": sys.version.split()[0], "sqlite": sqlite3.sqlite_version,
                    "openpyxl": openpyxl.__version__},
        "parametros": {"lote": LOTE, "transacao": "uma por aba", "foreign_keys": fk_ativo,
                       "ordem_carga": ORDEM_CARGA, "datas": "TEXT ISO 8601 (AAAA-MM-DDTHH:MM:SS)"},
        "ddl": {t: ddl(t) for t in ORDEM_CARGA},
        "indices_criados_apos_carga": [{"nome": n, "tabela": t, "colunas": c} for n, t, c in INDICES],
        "segundos_indices_e_analyze": seg_indices,
        "cargas": cargas,
        "contagens": contagens,
        "integrity_check": integrity,
        "fks_declaradas": fks_declaradas,
        "foreign_key_check_linhas": len(fk_check),
        "orfaos_por_relacao": orfaos,
        "nulos_por_coluna": nulos,
        "entrada_sha256_antes": hashes_antes,
        "entrada_sha256_depois": hashes_depois,
        "entrada_intacta": entrada_intacta,
        "gate": {"tabelas": len(tabelas_db), "falhas": falhas, "passou": not falhas},
        "segundos_total": round(time.perf_counter() - inicio, 3),
    }
    (SAIDA / "carga_sqlite.json").write_text(json.dumps(relatorio, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"tabelas: {len(tabelas_db)} | integrity_check: {integrity[0]} | foreign_key_check: {len(fk_check)} linhas "
          f"| FKs declaradas: {len(fks_declaradas)} | entrada intacta: {entrada_intacta}")
    if falhas:
        for f in falhas:
            print("FALHA:", f)
        raise SystemExit(1)
    print("GATE B PASSOU")


if __name__ == "__main__":
    main()

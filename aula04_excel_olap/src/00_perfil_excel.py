"""Fase A - perfil dos quatro Excels da Aula 04.

Lê entrada/*.xlsx com openpyxl em modo read_only (linha a linha, sem montar a
aba em memória), confere o SHA256 de cada arquivo contra o manifesto e grava
saida/perfil_excel.json com arquivo, aba, cabeçalhos, contagem de linhas e
amostra de tipos por coluna. Não escreve nada em entrada/.

Gate: quatro arquivos, nove abas e nove contagens iguais ao manifesto.
"""
import hashlib
import json
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
ENTRADA = BASE / "entrada"
SAIDA = BASE / "saida"
MANIFESTO = BASE / "contrato" / "manifesto_olist.json"
LIMITE_NOME_ABA = 31  # limite do próprio Excel para nome de aba
AMOSTRA = 3


def sha256(caminho):
    h = hashlib.sha256()
    with open(caminho, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()


def localizar_aba(wb, tabela):
    """Aba com o nome da tabela; o Excel corta nomes acima de 31 caracteres."""
    if tabela in wb.sheetnames:
        return tabela
    cortado = tabela[:LIMITE_NOME_ABA]
    if cortado in wb.sheetnames:
        return cortado
    raise SystemExit(f"FALHA: aba da tabela {tabela} não encontrada; abas: {wb.sheetnames}")


def perfilar_aba(ws):
    linhas = ws.iter_rows(values_only=True)
    cabecalho = list(next(linhas))
    tipos = [Counter() for _ in cabecalho]
    amostras = [[] for _ in cabecalho]
    n = 0
    celulas_faltantes = 0
    celulas_excedentes = 0
    largura = len(cabecalho)
    for linha in linhas:
        n += 1
        # Em read_only o openpyxl não completa células vazias no fim da linha:
        # a linha curta é completada com nulos; célula além do cabeçalho é contada.
        if len(linha) < largura:
            celulas_faltantes += largura - len(linha)
            linha = tuple(linha) + (None,) * (largura - len(linha))
        elif len(linha) > largura:
            celulas_excedentes += sum(1 for v in linha[largura:] if v is not None)
            linha = linha[:largura]
        for i, valor in enumerate(linha):
            tipos[i][type(valor).__name__ if valor is not None else "null"] += 1
            if valor is not None and len(amostras[i]) < AMOSTRA:
                amostras[i].append(valor if isinstance(valor, (int, float, str)) else str(valor))
    colunas = [
        {"nome": nome, "tipos_observados": dict(tipos[i]), "amostra": amostras[i]}
        for i, nome in enumerate(cabecalho)
    ]
    extras = {"celulas_vazias_no_fim_da_linha": celulas_faltantes,
              "celulas_preenchidas_fora_do_cabecalho": celulas_excedentes}
    return cabecalho, n, colunas, extras


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    inicio = time.perf_counter()
    manifesto = json.loads(MANIFESTO.read_text(encoding="utf-8"))
    SAIDA.mkdir(exist_ok=True)

    esperados = {w["file"] for w in manifesto["workbooks"]}
    encontrados = {p.name for p in ENTRADA.glob("*.xlsx")}
    falhas = []
    if esperados != encontrados:
        falhas.append(f"arquivos divergentes: esperados {sorted(esperados)}, encontrados {sorted(encontrados)}")

    arquivos = []
    for w in manifesto["workbooks"]:
        caminho = ENTRADA / w["file"]
        if not caminho.exists():
            continue
        t0 = time.perf_counter()
        hash_lido = sha256(caminho)
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        abas = []
        tabelas_manifesto = [s["table"] for s in w["sheets"]]
        nomes_abas_manifesto = {localizar_aba(wb, t) for t in tabelas_manifesto}
        extras = sorted(set(wb.sheetnames) - nomes_abas_manifesto)
        if extras:
            falhas.append(f"{w['file']}: abas fora do manifesto {extras}")
        for s in w["sheets"]:
            nome_aba = localizar_aba(wb, s["table"])
            print(f"perfilando {w['file']} / {nome_aba} ...", flush=True)
            cabecalho, n, colunas, extras_linha = perfilar_aba(wb[nome_aba])
            ok_contagem = n == s["rows"]
            ok_cab = all(c not in (None, "") for c in cabecalho) and len(set(cabecalho)) == len(cabecalho)
            chave_ok = all(k in cabecalho for k in s["candidate_key"])
            if not ok_contagem:
                falhas.append(f"{s['table']}: {n} linhas, manifesto {s['rows']}")
            if not ok_cab:
                falhas.append(f"{s['table']}: cabeçalho vazio ou repetido")
            if extras_linha["celulas_preenchidas_fora_do_cabecalho"]:
                falhas.append(f"{s['table']}: células preenchidas além do cabeçalho")
            if not chave_ok:
                falhas.append(f"{s['table']}: chave candidata {s['candidate_key']} ausente do cabeçalho")
            abas.append({
                "tabela": s["table"],
                "aba": nome_aba,
                "nome_aba_cortado_pelo_excel": nome_aba != s["table"],
                "cabecalhos": cabecalho,
                "linhas_dados": n,
                "linhas_manifesto": s["rows"],
                "contagem_confere": ok_contagem,
                "chave_candidata": s["candidate_key"],
                "leitura_linhas": extras_linha,
                "colunas": colunas,
            })
        wb.close()
        ok_hash = hash_lido == w["sha256"]
        if not ok_hash:
            falhas.append(f"{w['file']}: SHA256 {hash_lido} difere do manifesto")
        arquivos.append({
            "arquivo": w["file"],
            "bytes": caminho.stat().st_size,
            "sha256": hash_lido,
            "sha256_manifesto": w["sha256"],
            "sha256_confere": ok_hash,
            "abas": abas,
            "segundos": round(time.perf_counter() - t0, 3),
        })

    n_abas = sum(len(a["abas"]) for a in arquivos)
    n_abas_manifesto = sum(len(w["sheets"]) for w in manifesto["workbooks"])
    gate = {
        "arquivos": len(arquivos),
        "arquivos_manifesto": len(manifesto["workbooks"]),
        "abas": n_abas,
        "abas_manifesto": n_abas_manifesto,
        "contagens_conferem": sum(1 for a in arquivos for s in a["abas"] if s["contagem_confere"]),
        "falhas": falhas,
        "passou": not falhas and len(arquivos) == len(manifesto["workbooks"]) and n_abas == n_abas_manifesto,
    }
    perfil = {
        "gerado_em": datetime.now().astimezone().isoformat(timespec="seconds"),
        "script": "src/00_perfil_excel.py",
        "leitura": "openpyxl read_only=True, data_only=True, iter_rows(values_only=True)",
        "versoes": {"python": sys.version.split()[0], "openpyxl": openpyxl.__version__},
        "manifesto": {"arquivo": "contrato/manifesto_olist.json", "dataset": manifesto["dataset"],
                       "generated_on": manifesto["generated_on"]},
        "arquivos": arquivos,
        "gate": gate,
        "segundos_total": round(time.perf_counter() - inicio, 3),
    }
    destino = SAIDA / "perfil_excel.json"
    destino.write_text(json.dumps(perfil, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"gravado {destino.relative_to(BASE)}")
    print(f"gate: {gate['arquivos']} arquivos, {gate['abas']} abas, "
          f"{gate['contagens_conferem']}/{gate['abas_manifesto']} contagens iguais ao manifesto")
    if not gate["passou"]:
        for f in falhas:
            print("FALHA:", f)
        raise SystemExit(1)
    print("GATE A PASSOU")


if __name__ == "__main__":
    main()
